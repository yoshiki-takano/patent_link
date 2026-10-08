import re
from copy import copy
from dataclasses import dataclass
from io import BytesIO
from urllib.parse import quote

from openpyxl import load_workbook

ESPACENET_BASE = "https://worldwide.espacenet.com/patent/search?q=pn%3D"
JPLATPAT_BASE = "https://www.j-platpat.inpit.go.jp/c1801/PU"

_PATTERN = re.compile(r"^([A-Z]{2})(\d+)([A-Z]\d?)$")


@dataclass(frozen=True)
class PatentNumber:
    country: str
    number: str
    kind: str


def parse_number(raw) -> PatentNumber | None:
    if raw is None:
        return None
    text = re.sub(r"[\s/,\-]", "", str(raw)).upper()
    m = _PATTERN.match(text)
    if not m:
        return None
    return PatentNumber(*m.groups())


def _is_granted(p: PatentNumber) -> bool:
    return p.kind.startswith(("B", "C"))


def espacenet_number(p: PatentNumber) -> str:
    number = p.number
    if p.country == "JP" and _is_granted(p):
        number = number.lstrip("0")
    elif p.country == "KR" and _is_granted(p) and len(number) == 7:
        # Source data drops the "10" prefix of KR registration numbers.
        number = "10" + number
    elif p.country == "US" and p.kind.startswith("A") and len(number) == 11:
        # Espacenet (DOCDB) uses a 6-digit serial: US2016222235A1, not US20160222235A1.
        number = number[:4] + number[4:].lstrip("0").zfill(6)
    return f"{p.country}{number}{p.kind}"


def espacenet_url(p: PatentNumber | None) -> str | None:
    if p is None:
        return None
    return ESPACENET_BASE + quote(espacenet_number(p))


def _jplatpat_foreign_id(p: PatentNumber) -> str | None:
    kind, n = p.kind[0], p.number
    if p.country == "WO" and kind == "A" and len(n) == 10:
        return f"WO-A-{n[:4]}-{n[4:]}"
    if p.country == "US":
        if kind == "A" and len(n) == 11:
            return f"US-A-{n[:4]}-{n[4:]}"
        if kind == "B":
            return f"US-B-{n.zfill(9)}"
    if p.country == "EP" and kind in "AB":
        return f"EP-{kind}-{n.zfill(8)}"
    if p.country == "CN" and kind in "ABCUY":
        return f"CN-{kind}-{n}"
    if p.country == "KR":
        # J-PlatPat KR publication = 2-digit year + 7-digit serial.
        if kind == "A" and len(n) in (10, 11):
            return f"KR-A-{n[2:4]}{n[4:].zfill(7)}"
        if kind == "B":
            return f"KR-B-{('10' + n) if len(n) == 7 else n}"
    return None


def jplatpat_url(p: PatentNumber | None) -> str | None:
    if p is None:
        return None
    if p.country == "JP":
        if _is_granted(p):
            return f"{JPLATPAT_BASE}/JP-{p.number.lstrip('0').zfill(7)}/15/ja"
        if p.kind.startswith("A") and len(p.number) == 10:
            return f"{JPLATPAT_BASE}/JP-{p.number[:4]}-{p.number[4:]}/11/ja"
        return None
    doc_id = _jplatpat_foreign_id(p)
    return f"{JPLATPAT_BASE}/{doc_id}/50/ja" if doc_id else None


def build_links(values) -> list[tuple[str | None, str | None]]:
    links = []
    for v in values:
        p = parse_number(v)
        links.append((espacenet_url(p), jplatpat_url(p)))
    return links


def add_links_to_workbook(
    xlsx_bytes: bytes, sheet_name: str, column_name: str, header_row: int = 1
) -> bytes:
    wb = load_workbook(BytesIO(xlsx_bytes))
    ws = wb[sheet_name]

    header_cell = next(
        (c for c in ws[header_row] if c.value is not None and str(c.value).strip() == column_name),
        None,
    )
    if header_cell is None:
        raise ValueError(f"列「{column_name}」が見つかりません。")

    src_col = header_cell.column
    esp_col = ws.max_column + 1
    jpp_col = esp_col + 1

    for col, title in ((esp_col, "Espacenet"), (jpp_col, "J-PlatPat")):
        cell = ws.cell(row=header_row, column=col, value=title)
        font = copy(header_cell.font)
        font.name = "Meiryo UI"
        font.sz = 10
        font.scheme = None
        cell.font = font
        cell.fill = copy(header_cell.fill)
        cell.border = copy(header_cell.border)
        alignment = copy(header_cell.alignment)
        alignment.vertical = "top"
        cell.alignment = alignment
        ws.column_dimensions[cell.column_letter].width = 20

    for row in range(header_row + 1, ws.max_row + 1):
        p = parse_number(ws.cell(row=row, column=src_col).value)
        for col, url, label in (
            (esp_col, espacenet_url(p), "Espacenet で開く"),
            (jpp_col, jplatpat_url(p), "J-PlatPat で開く"),
        ):
            cell = ws.cell(row=row, column=col, value=label if url else None)
            if url:
                cell.hyperlink = url
                cell.style = "Hyperlink"
            font = copy(cell.font)
            font.name = "Meiryo UI"
            font.sz = 10
            font.scheme = None
            cell.font = font
            alignment = copy(cell.alignment)
            alignment.vertical = "top"
            cell.alignment = alignment

    out = BytesIO()
    wb.save(out)
    return out.getvalue()

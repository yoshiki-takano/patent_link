from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st

from openpyxl import Workbook, load_workbook

from patent_links import JPLATPAT_NOTE, add_links_to_workbook, build_links, jplatpat_may_be_missing, parse_number

st.set_page_config(page_title="公報番号リンク作成", page_icon=":material/link:", layout="wide")


@st.cache_data(ttl="1h", scope="session")
def read_sheets(xlsx_bytes: bytes) -> dict[str, list[list[str | None]]]:
    # Read raw rows (no header) so row numbers match Excel exactly.
    wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    sheets = {
        ws.title: [[None if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        for ws in wb.worksheets
    }
    wb.close()
    return sheets


@st.cache_data(ttl="1h", scope="session")
def build_output(
    xlsx_bytes: bytes,
    sheet: str,
    src_col: int,
    header_row: int,
    link_rules_signature: str,
) -> bytes:
    return add_links_to_workbook(xlsx_bytes, sheet, src_col, header_row)


@st.cache_data(ttl="1h", scope="session")
def build_text_output(numbers: tuple[str, ...], link_rules_signature: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "公報番号"
    ws.cell(row=1, column=1, value="公報番号")
    for row, number in enumerate(numbers, start=2):
        ws.cell(row=row, column=1, value=number)

    xlsx_buffer = BytesIO()
    wb.save(xlsx_buffer)
    wb.close()
    return add_links_to_workbook(xlsx_buffer.getvalue(), ws.title, 1, 1)


def show_result(values: pd.Series, column: str) -> None:
    links = build_links(values)
    notes = [JPLATPAT_NOTE if jplatpat_may_be_missing(parse_number(v)) else None for v in values]
    result = pd.DataFrame(
        {
            column: values,
            "Espacenet": [espacenet for espacenet, _ in links],
            "J-PlatPat": [jplatpat for _, jplatpat in links],
            "J-PlatPat 注意": notes,
        }
    )

    total = int(values.notna().sum())
    unparsed = int(sum(v is not None and not pd.isna(v) and parse_number(v) is None for v in values))
    no_jpp = int(result["J-PlatPat"].isna().sum() - values.isna().sum() - unparsed)

    m1, m2, m3 = st.columns(3)
    m1.metric("公報番号", total, border=True)
    m2.metric("解析失敗", unparsed, border=True)
    m3.metric("J-PlatPat 未対応", no_jpp, border=True)

    flagged = [str(v) for v, note in zip(values, notes) if note]
    if flagged:
        st.warning(
            "次の番号は、J-PlatPatの蓄積範囲外の可能性があります。\n\n" + "、".join(flagged),
            icon=":material/warning:",
        )

    st.dataframe(
        result,
        column_config={
            "Espacenet": st.column_config.LinkColumn("Espacenet", display_text="Espacenet で開く"),
            "J-PlatPat": st.column_config.LinkColumn("J-PlatPat", display_text="J-PlatPat で開く"),
        },
        hide_index=True,
    )


def detect_header_row(rows: list[list[str | None]], max_rows: int = 20) -> int:
    for i, row in enumerate(rows[:max_rows], start=1):
        if any(v is not None and v.strip() == "公報番号" for v in row):
            return i
    return 1


st.title(":material/link: 公報番号リンク作成")
st.caption("Derwent Patent Searchの公報番号から Espacenet / J-PlatPat へのリンクを作成します。ただし、欠番、未蓄積、落丁の場合があります。")

st.markdown(
    """
**リンク作成ルール**

| リンク先 | 対応範囲・番号の扱い |
|---|---|
| Espacenet | 公報番号として解析できる国コード付き番号を検索します。日本の登録公報は先頭の0を除去します（例: JP02500001Y2 → JP2500001Y2）。先頭が0以外の7桁B1/B2は特公平形式に変換し、先頭が`1`で通番が`12440`以下なら昭和64年、それ以外は平成元年以降にします（例: JP1012440B2 → JPS6412440B2、JP1012441B2 → JPH0112441B2）。8桁の特公昭形式は元号・年に変換します（例: JP46012531B2 → JPS4612531B2）。中国の10桁A/B公報は先頭2桁を除去（例: CN1988102056A → CN88102056A、CN1985100001B → CN85100001B）、韓国の登録公報は6桁番号が186214以下なら`0`、186215以上なら`100`、7桁番号は`10`を前に付加（例: KR107409B1 → KR0107409B1、KR990264B1 → KR100990264B1、KR13B1 → KR1000013B1）、9桁番号は西暦4桁の年を2桁にし通番を7桁に補正（例: KR199513437B1 → KR950013437B1）、6桁Y1公報は140846以下なら`0`、140847以上なら`200`を前に付加（例: KR106898Y1 → KR0106898Y1、KR485463Y1 → KR200485463Y1）、9桁Y1公報は西暦4桁の年を2桁にし通番を7桁に補正（例: KR199600937Y1 → KR960000937Y1）、権利種別コード`20`付き13桁U公報は2桁年＋7桁通番に補正（例: KR2083000003695U → KR830003695U）、米国の公開公報は年の後の先頭0を除去します。検索結果はEspacenetの収録状況に依存します。 |
| J-PlatPat（日本） | 公開・公表公報（A）は種別`11`、通常の特許登録公報（B/B1/B2）と特許発明明細書（C、先頭の0を除去。例: JP0000001C → JP-1/15/ja）は`15`、登録実用新案公報（U）・実用新案登録公報（Y、0始まりの8桁）は`25`で照会します（例: JP02500001Y2 → JP-2500001/25/ja）。特公昭の8桁および先頭が0以外の7桁B1/B2は、元号・年・6桁通番の形式で種別`12`を指定します（例: JP46012531B2 → JP-S46-012531/12/ja、JP7013126B1 → JP-H07-013126/12/ja）。7桁B1/B2のうちJP10から始まる番号は、J-PlatPatではすべてH01にします。 |
| J-PlatPat（外国） | WO（A）、US（A/B。7桁以下のA（旧特許）は9桁に0埋め（例: US999999A → US-A-000999999））、EP（A/B。番号は9桁に0埋め（例: EP6042A1 → EP-A-000006042、EP3086489A3 → EP-A-003086489、EP1234567B1 → EP-B-001234567））、CN（A/B/C/U/Y。10桁の年形式A/Bは西暦の上2桁を除き9桁に0埋め（例: CN1985100001B → CN-B-085100001）、7桁番号は9桁に0埋め（例: CN1536611A → CN-A-001536611、CN1019873C → CN-C-001019873、CN2891917Y → CN-Y-002891917））、KR（A/B/U/Y。KR-Bは7桁以下を9桁に0埋め（例: KR990264B1 → KR-B-000990264、KR13B1 → KR-B-000000013）、KR199513437B1 → KR-B-950013437、KR-Uは例: KR201405583U → KR-U-140005583、KR830003893U → KR-U-830003893、KR2083000003695U → KR-U-830003695、KR-Y1は例: KR485463Y1 → KR-Y-000485463）、GB・DE・DD・FR・CH・CAにリンクを作成します。外国公報の種別コードは`50`です。 |
| J-PlatPat（対象外） | 上記以外の国・種別（例: TW）はリンク欄を空欄にします。 |
| J-PlatPat（注意） | J-PlatPatの文献蓄積情報（2026年10月時点）の最古の番号より古い番号や、蓄積されていない種別・形式（例: KR199600937Y1、GBのB公報）は、「J-PlatPat 注意」列に「蓄積範囲外の可能性」と表示します。リンクは作成されます。 |

リンクは各サービスの照会画面を開きます。対象文献がデータベースに収録されていない場合、検索結果が表示されないことがあります。
"""
)

input_mode = st.segmented_control("入力方法", ["Excelファイル", "直接入力"], default="Excelファイル")
link_rules_signature = sha256(Path(__file__).with_name("patent_links.py").read_bytes()).hexdigest()

if input_mode == "直接入力":
    raw_numbers = st.text_area(
        "DPS公報番号（1行に1件）",
        placeholder="JP2020123456A\nUS10123456B2",
        height=200,
    )
    numbers = tuple(line.strip() for line in raw_numbers.splitlines() if line.strip())
    if not numbers:
        st.info("公報番号を1行に1件ずつ入力してください。")
        st.stop()

    values = pd.Series(numbers, dtype=object)
    show_result(values, "公報番号")
    output = build_text_output(numbers, link_rules_signature)
    st.download_button(
        "リンク付きExcelをダウンロード",
        data=output,
        file_name="公報番号_with_links.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        icon=":material/download:",
        type="primary",
    )
else:
    uploaded = st.file_uploader("Excelファイル (.xlsx)", type=["xlsx"])
    if uploaded is None:
        st.stop()

    data = uploaded.getvalue()
    sheets = read_sheets(data)

    c1, c2, c3 = st.columns(3)
    sheet = c1.selectbox("シート", list(sheets))
    rows = sheets[sheet]
    if not rows:
        st.warning("選択したシートにデータがありません。")
        st.stop()

    header_row = c2.number_input(
        "見出し行（Excelの行番号）",
        min_value=1,
        max_value=len(rows),
        value=detect_header_row(rows),
        key=f"header_row_{sheet}",
        help="「公報番号」を含む行を自動で初期選択します。",
    )
    width = max(len(r) for r in rows)
    headers = rows[header_row - 1] + [None] * (width - len(rows[header_row - 1]))
    labels = [h.strip() if h and h.strip() else f"（無題）{i + 1}列目" for i, h in enumerate(headers)]
    col_idx = c3.selectbox(
        "公報番号列",
        range(width),
        index=labels.index("公報番号") if "公報番号" in labels else 0,
        format_func=lambda i: labels[i],
        key=f"column_{sheet}_{header_row}",
    )
    column = labels[col_idx]

    values = pd.Series(
        [r[col_idx] if col_idx < len(r) else None for r in rows[header_row:]], dtype=object
    )
    show_result(values, column)

    try:
        output = build_output(data, sheet, col_idx + 1, header_row, link_rules_signature)
    except ValueError as e:
        # Streamlit Cloud redacts exception text; openpyxl messages only contain row/column numbers.
        st.error(
            f"リンク付きExcelの作成に失敗しました（シート: {sheet}, 見出し行: {header_row!r}, "
            f"列: {col_idx + 1!r}, 列数: {width}, 行数: {len(rows)}）: {e}"
        )
        st.stop()

    st.download_button(
        "リンク付きExcelをダウンロード",
        data=output,
        file_name=f"{Path(uploaded.name).stem}_with_links.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        icon=":material/download:",
        type="primary",
    )
